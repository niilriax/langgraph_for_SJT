"""Add the planned structural-validity formula and Mussel MTMM method to a report."""

from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "15313A"
TEAL = "008277"
PALE = "F4F7F8"
BORDER = "D9D9D9"


def _font_path(bold: bool = False) -> str:
    candidates = (
        [r"C:\Windows\Fonts\msyhbd.ttc", r"C:\Windows\Fonts\simhei.ttf"]
        if bold
        else [r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simsun.ttc"]
    )
    for candidate in candidates:
        if Path(candidate).is_file():
            return candidate
    raise FileNotFoundError("找不到可用于公式图的中文字体")


def make_formula_card(output_path: Path) -> None:
    width, height = 1600, 610
    canvas = Image.new("RGB", (width, height), "#FAF9F5")
    draw = ImageDraw.Draw(canvas)
    title_font = ImageFont.truetype(_font_path(True), 43)
    body_font = ImageFont.truetype(_font_path(False), 29)
    lead_font = ImageFont.truetype(_font_path(True), 29)
    equation_font = ImageFont.truetype(r"C:\Windows\Fonts\cambria.ttc", 71)

    draw.text((18, 16), "6. 结构效度：五特质 × 两方法 MTMM", font=title_font, fill="#062E45")
    box = (18, 88, width - 18, 442)
    draw.rounded_rectangle(box, radius=34, fill=f"#{NAVY}")

    equations = (
        "Yₜₘₚ = λₜ · Tₜ + λₘ · Mₘ + εₜₘₚ",
        "Σ(θ) = Λₜ Φₜ Λₜᵀ + Λₘ Φₘ Λₘᵀ + Θ",
    )
    equation_y = (165, 290)
    for text, y in zip(equations, equation_y):
        bounds = draw.textbbox((0, 0), text, font=equation_font)
        text_width = bounds[2] - bounds[0]
        draw.text(((width - text_width) / 2, y), text, font=equation_font, fill="#FFFFFF")

    lead = "说明："
    explanation = "检验五个 facet 的共同变异主要来自特质，而不是测量方法或单一总因子。"
    draw.text((18, 466), lead, font=lead_font, fill=f"#{TEAL}")
    lead_width = draw.textlength(lead, font=lead_font)
    draw.text((18 + lead_width + 6, 466), explanation, font=body_font, fill="#20323A")
    draw.text(
        (18, 526),
        "当前不填 A/B/C 数值；待五 facet 问卷及对应效标数据齐备后再实施。",
        font=body_font,
        fill="#4E5963",
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output_path, dpi=(300, 300))


def _set_run_font(run, size: float = 11.0, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _format_body_paragraph(paragraph) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    for run in paragraph.runs:
        _set_run_font(run, 11.0)


def _add_body(document: Document, lead: str, text: str):
    paragraph = document.add_paragraph(style="Normal")
    lead_run = paragraph.add_run(lead)
    _set_run_font(lead_run, 11.0, bold=True)
    body_run = paragraph.add_run(text)
    _set_run_font(body_run, 11.0)
    _format_body_paragraph(paragraph)
    return paragraph


def _move_after(anchor, element) -> None:
    anchor.addnext(element)


def _set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)


def _format_mtmm_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.55), Inches(1.02), Inches(0.86), Inches(0.96), Inches(0.96)]
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.width = widths[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            if row_index == 0:
                _set_cell_fill(cell, NAVY)
            elif row_index % 2 == 0:
                _set_cell_fill(cell, PALE)
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    _set_run_font(
                        run,
                        9.0,
                        bold=row_index == 0,
                        color="FFFFFF" if row_index == 0 else "202A30",
                    )
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_tr_pr.append(repeat)
    _set_table_borders(table)


def _set_alt_text(paragraph, description: str) -> None:
    doc_props = paragraph._p.xpath(".//wp:docPr")
    if doc_props:
        doc_props[0].set("descr", description)


def _replace_appendix_symbols(document: Document) -> None:
    symbol_table = None
    for table in document.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == "符号":
            symbol_table = table
            break
    if symbol_table is None:
        return
    replacements = [
        ("Y_tmp", "特质 t、方法 m、题包 p 对应的观测题包得分"),
        ("λ_T / λ_M", "观测题包在特质因子／方法因子上的载荷"),
        ("T_t / M_m / ε_tmp", "潜在特质因子／潜在方法因子／题包残差"),
    ]
    for row, (symbol, meaning) in zip(symbol_table.rows[-3:], replacements):
        row.cells[0].text = symbol
        row.cells[1].text = meaning
    extra = symbol_table.add_row()
    extra.cells[0].text = "Λ, Φ, Θ"
    extra.cells[1].text = "载荷矩阵、潜变量协方差矩阵和残差协方差矩阵"
    for row_index, row in enumerate(symbol_table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    _set_run_font(run, 9.5, bold=row_index == 0, color="FFFFFF" if row_index == 0 else "202A30")
            if row_index == 0:
                _set_cell_fill(cell, NAVY)
            elif row_index % 2 == 0:
                _set_cell_fill(cell, PALE)
    _set_table_borders(symbol_table)


def edit_document(source: Path, output: Path, formula_card: Path) -> None:
    document = Document(source)

    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("4  整卷五项指标及公式解释"):
            paragraph.text = "4  整卷指标及公式解释"
            break

    paragraphs = list(document.paragraphs)
    structure_heading = next(
        paragraph for paragraph in paragraphs if paragraph.text.strip().startswith("4.6  结构效度")
    )
    next_heading = next(
        paragraph for paragraph in paragraphs if paragraph.text.strip().startswith("5  A/B/C")
    )
    structure_heading.text = "4.6  结构效度与 Mussel 的五 facet 验证方法"

    body = document._element.body
    active = False
    to_remove = []
    for child in list(body):
        if child is structure_heading._p:
            active = True
            continue
        if child is next_heading._p:
            break
        if active:
            to_remove.append(child)
    for child in to_remove:
        body.remove(child)

    anchor = structure_heading._p

    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(4)
    image_paragraph.paragraph_format.space_after = Pt(4)
    image_paragraph.add_run().add_picture(str(formula_card), width=Inches(5.55))
    _set_alt_text(
        image_paragraph,
        "五特质两方法多特质多方法模型公式，观测题包同时由特质因子、方法因子和残差解释。",
    )
    _move_after(anchor, image_paragraph._p)
    anchor = image_paragraph._p

    caption = document.add_paragraph(style="Normal")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption.add_run("图 2  五特质 × 两方法 MTMM 结构效度模型。")
    _set_run_font(caption_run, 9.5, color="4E5963")
    caption.paragraph_format.space_after = Pt(8)
    _move_after(anchor, caption._p)
    anchor = caption._p

    paragraphs_to_add = [
        (
            "指标含义：",
            "结构效度检验五个目标 facet 是否形成彼此可区分的潜在特质结构，以及这种结构是否主要来自特质本身，而不是来自 SJT 或自陈问卷的测量方法。这里的 Y_tmp 表示特质 t、方法 m、题包 p 的观测得分；λ_T 和 λ_M 分别表示特质载荷与方法载荷。",
        ),
        (
            "Mussel 的数据组织：",
            "研究同时使用五个 SJT facet 和五个对应的 NEO-PI-R facet。每个“facet × 方法”组合被划分为 3 个题包：SJT 每个 facet 的 22 题随机分入 3 包后取平均，NEO-PI-R 每个 facet 的 8 题也分为 3 包。因此模型包含 5 个特质 × 2 种方法 × 3 个题包，共 30 个观测变量。",
        ),
        (
            "Mussel 的核心模型：",
            "使用最大似然法拟合相关特质相关方法模型（CTCM）：5 个相关特质因子表示 N4、E2、O5、A4 和 C5，2 个相关方法因子表示 SJT 与 NEO-PI-R。研究随后将该模型与仅特质模型、仅方法模型和单一总因子模型比较，并分别对 SJT 与 NEO-PI-R 拟合五特质模型。",
        ),
    ]
    for lead, text in paragraphs_to_add:
        paragraph = _add_body(document, lead, text)
        _move_after(anchor, paragraph._p)
        anchor = paragraph._p

    table_data = [
        ["模型", "χ²/df", "CFI", "RMSR", "RMSEA"],
        ["Model 1  特质＋方法 CTCM", "1.51", ".93", ".04", ".05"],
        ["Model 2  仅相关特质 CT", "1.81", ".88", ".05", ".06"],
        ["Model 3  仅相关方法 CM", "5.59", ".30", ".13", ".14"],
        ["Model 4  单一总因子 1F", "5.61", ".29", ".13", ".14"],
        ["Model 5  仅 SJT 五特质", "1.10", ".99", ".00", ".02"],
        ["Model 6  仅 NEO 五特质", "2.29", ".92", ".09", ".07"],
    ]
    table = document.add_table(rows=len(table_data), cols=len(table_data[0]))
    for row_index, row in enumerate(table_data):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    _format_mtmm_table(table)
    _move_after(anchor, table._tbl)
    anchor = table._tbl

    concluding = [
        (
            "Mussel 的结果解释：",
            "完整 CTCM 模型达到可接受拟合；潜在特质解释 40.6% 的方差，潜在方法因子解释 6.0%。仅特质模型明显优于仅方法模型和单因子模型；SJT 自身的五特质模型拟合也优于 NEO-PI-R 五特质模型。作者据此认为，SJT 的作答差异主要反映五个可区分的目标特质，而不是笼统的方法效应或单一总因子。",
        ),
        (
            "本项目的后续方案：",
            "当前 A/B/C 实验只覆盖单个目标 facet，因此本阶段不填报结构效度数值。待 N4、E2、O5、A4、C5 五套问卷及相应 IPIP-NEO facet 数据齐备后，再按五特质模型检验。为了与 Mussel 对照，可报告题包层面的 CTCM；为了减少随机题包掩盖单题失配的风险，还应补充题目层面的五因子有序 CFA。",
        ),
        (
            "报告指标：",
            "至少报告 χ²、df、χ²/df、CFI、TLI、RMSEA（含区间）和 SRMR/RMSR，并比较五特质模型与单因子模型。结构效度只有在完整五 facet 数据收集后才能形成证据，不能根据当前单 facet 结果提前下结论。",
        ),
        (
            "文献来源：",
            "Mussel, P., Gatzka, T., & Hewig, J. (2018). Situational Judgment Tests as an Alternative Measure for Personality Assessment. European Journal of Psychological Assessment, 34(5), 328–335. https://doi.org/10.1027/1015-5759/a000346",
        ),
    ]
    for lead, text in concluding:
        paragraph = _add_body(document, lead, text)
        _move_after(anchor, paragraph._p)
        anchor = paragraph._p

    _replace_appendix_symbols(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--formula-card", type=Path, required=True)
    args = parser.parse_args()
    make_formula_card(args.formula_card)
    edit_document(args.source, args.output, args.formula_card)
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
